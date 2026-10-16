#!/usr/bin/env python3
"""Evidence ledger; semantic judgments are supplied by the invoking Agent."""
import argparse, copy, json, os, tempfile
from pathlib import Path
from datetime import datetime, timezone, timedelta
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
DEFAULTS=json.loads((Path(__file__).resolve().parents[1]/'references/source-presets.json').read_text())['sources']
def search_plan(s):
 policy=s.get('search_policy',{})
 return {'priority_groups':policy.get('priority_groups',[]),'allow_expansion':policy.get('allow_expansion',True),'sources':{k:v for k,v in {**DEFAULTS,**s['source_overrides']}.items() if v.get('enabled')},'note':'Priority controls search coverage, not personal fit or eligibility; log actual searches separately.'}
APPLICATIONS={'not_applied','saved','applied','interviewing','offer','rejected','dismissed','withdrawn'}
HEADER_ZH={'编号':'id','职位链接':'url','岗位链接':'url','公司':'company','岗位':'title','投递状态':'application_status'}
APPLICATION_ZH={'未投递':'not_applied','已收藏':'saved','已投递':'applied','面试中':'interviewing','已获Offer':'offer','被拒绝':'rejected','不再考虑':'dismissed','已撤回':'withdrawn'}
REL={'Directly Relevant':0,'Adjacent':1,'Exploratory':2,'Not Relevant':3}
FIT={'Strong':0,'Partial':1,'Needs clarification':2}
def now(): return datetime.now(timezone.utc).isoformat()
def urlkey(url):
 p=urlsplit(url)
 if p.scheme not in {'https','http'} or not p.netloc: raise ValueError('Expected observed HTTP(S) URL')
 return urlunsplit((p.scheme.lower(),p.netloc.lower(),p.path.rstrip('/'),urlencode([(k,v) for k,v in parse_qsl(p.query) if not k.startswith('utm_') and k not in {'source','trk'}]),''))
def empty(): return {'schema_version':1,'next_id':1,'jobs':[],'profile':{},'source_overrides':{},'questions':{},'log':[],'conflicts':[]}
def load(path):
 s=json.loads(Path(path).read_text()) if Path(path).exists() else empty()
 if s.get('schema_version')!=1: raise ValueError('Unsupported state version')
 return s
def mode(s):
 p=s['profile'];return 'Matching' if p.get('matching_ready') and any(x.get('evidence') and x.get('summary') for x in p.get('experience',[])) else 'Discovery'
REQUIRED_INTAKE=('target','locations','experience_years','education')
OPTIONAL_INTAKE=('salary','industries','company_preferences','work_arrangement','employment_type','skills','achievements','resume','availability','exclusions')
def intake(s):
 p=s['profile'];missing=[]
 for key in REQUIRED_INTAKE:
  value=p.get(key)
  if key=='experience_years':
   valid=isinstance(value,(int,float)) and not isinstance(value,bool) and value>=0
  elif key=='locations':valid=isinstance(value,list) and any(isinstance(x,str) and x.strip() for x in value)
  else:valid=isinstance(value,str) and bool(value.strip()) and value.strip().lower() not in {'unknown','未知','待确认'}
  if not valid:missing.append(key)
 return {'ready':not missing,'missing_required':missing,'optional_to_offer':[k for k in OPTIONAL_INTAKE if k not in p], 'action':'Proceed' if not missing else 'Ask required before formal search'}

def coverage(s):
 records=s.get('search_log',[])
 kinds={r['kind'] for r in records if r.get('outcome')=='searched'}
 return {'sources_attempted':len({r['source_id'] for r in records}), 'channels_searched':sorted(kinds), 'unique_jobs':len(s['jobs']), 'recommended':len(recommend(s)), 'official_confirmed_open':sum(status(j)==('open','Verified') for j in s['jobs']), 'unverified_or_conflicting':sum(status(j)[1]!='Verified' for j in s['jobs']), 'partial_jds':sum(j.get('jd_completeness')!='full' for j in s['jobs']), 'search_log':records, 'warnings':(['Both official and third-party channels need search or explicit unavailable evidence'] if kinds!={'official','third_party'} else [])+(['Fewer than Top10: report coverage, exclusions and stopping reason; do not imply market exhaustion'] if len(recommend(s))<10 else [])}

def status(j):
 latest={}
 for o in j.get('observations',[]):
  k=urlkey(o['url'])
  if k not in latest or datetime.fromisoformat(o['checked_at'])>datetime.fromisoformat(latest[k]['checked_at']): latest[k]=o
 recent=[o for o in latest.values() if datetime.fromisoformat(o['checked_at'])>=datetime.now(timezone.utc)-timedelta(days=7)]
 official=[o for o in recent if o['kind']=='official']; third=[o for o in recent if o['kind']=='third_party']
 chosen=official or third
 outcomes={o['outcome'] for o in chosen if o['outcome']!='unknown'}
 current=next(iter(outcomes)) if len(outcomes)==1 else 'unknown'
 # A latest failed official check must not resurrect stale open third-party evidence.
 if official and any(o['outcome']=='unknown' for o in official): current='unknown'
 conflict=len({o['outcome'] for o in recent if o['outcome']!='unknown'})>1
 verification='Conflict' if conflict else ('Verified' if official and current!='unknown' else 'Unverified')
 return current,verification

def eligible(s,j):
 if j['application_status'] in {'rejected','dismissed','withdrawn'} or status(j)[0]=='closed' or j['relevance']=='Not Relevant':return False
 if mode(s)=='Matching' and any(r.get('hard_exclusion') and r.get('assessment')=='Not Met' for r in j.get('requirements',[])):return False
 return True

def recommend(s,sort='relevance',limit=10,include_unverified=False):
 if sort=='match' and mode(s)!='Matching': raise ValueError('Matching needs referenced experience; use relevance')
 rows=[copy.deepcopy(j) for j in s['jobs'] if eligible(s,j) and (include_unverified or any(o.get('evidence') and o.get('url') for o in j.get('observations',[])))]
 for j in rows:
  j['mode']=mode(s);j['recruiting_status'],j['verification_status']=status(j)
  if mode(s)=='Discovery': j.pop('fit',None);j.pop('fit_reason',None)
 rows.sort(key=lambda j:j['id'])
 if sort=='latest':rows.sort(key=lambda j:j.get('published_at') or '',reverse=True)
 elif sort=='match':rows.sort(key=lambda j:(FIT.get(j.get('fit'),2),REL[j['relevance']]))
 else:rows.sort(key=lambda j:REL[j['relevance']])
 return rows[:limit] if limit else rows

def validate_job(j):
 for k in ('company','title','location','url','relevance','relevance_reason'):
  if not j.get(k):raise ValueError('Missing job '+k)
 urlkey(j['url'])
 if j['relevance'] not in REL:raise ValueError('Invalid relevance')
 if j.get('fit') and (j['fit'] not in FIT or not j.get('fit_reason')):raise ValueError('Fit requires supported label and reason')
 if j.get('jd_completeness','partial') not in {'full','partial','unavailable'}:raise ValueError('Invalid JD completeness')
 if j.get('fit')=='Strong' and j.get('jd_completeness')!='full':raise ValueError('Strong overall fit requires a fully read qualification section')
 if j.get('published_at'):
  datetime.fromisoformat(j['published_at'])
  if not j.get('published_evidence'):raise ValueError('Publication date needs source evidence')
 for o in j.get('observations',[]):
  urlkey(o['url']);d=datetime.fromisoformat(o['checked_at'])
  if not d.tzinfo or d>datetime.now(timezone.utc)+timedelta(minutes=5):raise ValueError('Invalid observation time')
  if o['kind'] not in {'official','third_party'} or o['outcome'] not in {'open','closed','unknown'} or not o.get('evidence') or not o.get('method'):raise ValueError('Invalid observation')
 for r in j.get('requirements',[]):
  if not all(r.get(k) for k in ('key','kind','criticality','evidence','url')):raise ValueError('Incomplete requirement evidence')
  urlkey(r['url'])
  if r['kind'] not in {'skill','education','experience'}:raise ValueError('Invalid requirement kind; verification limitations are not candidate requirements')
  if r['criticality'] not in {'Mandatory','Preferred','Contextual'} or r.get('assessment','Unknown') not in {'Met','Unknown','Not Met'}:raise ValueError('Invalid assessment')
  if r.get('assessment') in {'Met','Not Met'} and not r.get('candidate_evidence'):raise ValueError('Candidate assessment needs evidence')
  if r.get('hard_exclusion'):
   if r.get('assessment')!='Not Met' or r['criticality']!='Mandatory':raise ValueError('Unknown/preferred cannot hard exclude')
   if r['kind']=='education' and not any(o['kind']=='official' and urlkey(o['url'])==urlkey(r['url']) for o in j.get('observations',[])):raise ValueError('Education exclusion requires official observation')

def ingest(s,b):
 if 'profile' in b:
  p=b['profile']
  for k,v in p.get('skills',{}).items():
   if v.get('state') not in {'Met','Unknown','Not Met'}:raise ValueError('Invalid skill '+k)
   if v['state']!='Unknown' and not v.get('evidence'):raise ValueError('Skill needs evidence')
  s['profile']=p
 for incoming in b.get('jobs',[]):
  validate_job(incoming);j=copy.deepcopy(incoming);j.setdefault('jd_completeness','partial')
  if 'jd_limitation' in j and 'jd_limitations' not in j:j['jd_limitations']=j.pop('jd_limitation')
  found=None
  for old in s['jobs']:
   same_req=j.get('requisition_id') and j['company'].casefold()==old['company'].casefold() and j.get('requisition_id')==old.get('requisition_id')
   same_url=urlkey(j['url']) in {urlkey(u) for u in old['urls']+[old['url']]}
   if j.get('merge_into')==old['id'] or same_req or same_url: found=old;break
  if j.get('merge_into') and not found:raise ValueError('Unknown merge target')
  if found:
   observations=list(found.get('observations',[]));observed={json.dumps(o,sort_keys=True,ensure_ascii=False) for o in observations}
   for o in j.get('observations',[]):
    signature=json.dumps(o,sort_keys=True,ensure_ascii=False)
    if signature not in observed:observations.append(o);observed.add(signature)
   urls=list(dict.fromkeys(found['urls']+[j['url']]))
   # User states and IDs are never ingested from a JD batch.
   keep={k:found[k] for k in ('id','application_status','first_seen','status_history')}
   found.update({k:v for k,v in j.items() if k not in {'id','application_status','status_history','first_seen','urls'}})
   found.update(keep);found['observations']=observations;found['urls']=urls;found['last_seen']=now()
   official_urls=[o['url'] for o in observations if o['kind']=='official']
   if official_urls:
    found['url']=official_urls[-1]
    found['urls']=list(dict.fromkeys(found['urls']+[found['url']]))
  else:
   j.update(id=f"J{s['next_id']:03d}",application_status='not_applied',first_seen=now(),last_seen=now(),urls=[j['url']],status_history=[])
   s['next_id']+=1;s['jobs'].append(j)
  logged={json.dumps(o,sort_keys=True,ensure_ascii=False) for o in s['log']}
  for o in j.get('observations',[]):
   entry={'job_id':found['id'] if found else j['id'],**o};signature=json.dumps(entry,sort_keys=True,ensure_ascii=False)
   if signature not in logged:s['log'].append(entry);logged.add(signature)
 return {'mode':mode(s),'jobs':len(s['jobs'])}

def clarify(s):
 jobs=[j for j in s['jobs'] if eligible(s,j)];n=len(jobs);groups={}
 for j in jobs:
  seen=set()
  for r in j.get('requirements',[]):
   k=r['key']
   if k in seen:continue
   seen.add(k);g=groups.setdefault(k,{'key':k,'count':0,'mandatory':0,'impact':0,'unknown_count':0,'unknown_mandatory':0,'unknown_impact':0})
   g['count']+=1;g['mandatory']+=r['criticality']=='Mandatory';g['impact']+=bool(r.get('decision_impact'))
   known=r.get('assessment') in {'Met','Not Met'}
   # Per-JD assessed evidence matters: a general skill may be known while a required level is unknown.
   if r.get('assessment') is None:known=s['profile'].get('skills',{}).get(k,{}).get('state','Unknown')!='Unknown'
   if not known:
    g['unknown_count']+=1;g['unknown_mandatory']+=r['criticality']=='Mandatory';g['unknown_impact']+=bool(r.get('decision_impact'))
 candidates=[]
 for k,g in groups.items():
  if not g['unknown_count'] or k in s['questions']:continue
  g['prevalence']=g['count']/n if n else 0
  g['action']='Ask' if g['unknown_mandatory']>=max(2,n*.2) and g['unknown_impact']>=max(2,n*.15) else 'Keep Unknown'
  candidates.append(g)
 candidates.sort(key=lambda g:(g['action']=='Ask',g['impact'],g['mandatory']),reverse=True)
 if s['profile'].get('broad_discovery'):action='Proceed'
 elif any(g['action']=='Ask' for g in candidates):action='Ask'
 elif n>=100:action='Refine'
 else:action='Proceed'
 return {'action':action,'relevant_jobs':n,'questions':[g for g in candidates if g['action']=='Ask'][:3] if action=='Ask' else [],'unknowns':candidates,'note':'Refine proposes evidence-based clusters; no filters applied before user answers.'}

def update_status(s,id,value):
 if value not in APPLICATIONS:raise ValueError('Invalid application status')
 j=next(j for j in s['jobs'] if j['id']==id)
 j['status_history'].append({'from':j['application_status'],'to':value,'at':now()});j['application_status']=value

def import_tracker(s,rows):
 pending=[]
 for r in rows:
  r={HEADER_ZH.get(k,k):v for k,v in r.items()}
  value=r.get('application_status');value=APPLICATION_ZH.get(value,value)
  if value not in APPLICATIONS:raise ValueError('Invalid tracker status')
  # IDs alone from another workspace are unsafe. Require a known observed URL.
  j=next((j for j in s['jobs'] if r.get('url') and urlkey(r['url']) in {urlkey(u) for u in j['urls']}),None)
  if not j:pending.append({'reason':'unmatched_url','row':r});continue
  if value!=j['application_status']:
   c={'id':f'C{len(s["conflicts"])+1:03d}','job_id':j['id'],'local':j['application_status'],'incoming':value,'resolved':False};s['conflicts'].append(c);pending.append(c)
 return pending

def export_data(s):
 rows=[]
 for j in s['jobs']:
  row={k:j.get(k) for k in ('id','company','title','location','url','relevance','application_status','published_at','relevance_reason','fit','fit_reason','first_seen','jd_completeness','jd_limitations')}
  row['recruiting_status'],row['verification_status']=status(j)
  row={k:row[k] for k in ('id','company','title','location','recruiting_status','verification_status','application_status','relevance','url','published_at','relevance_reason','fit','fit_reason','first_seen','jd_completeness','jd_limitations')};row['mode']=mode(s);row['requirements']=json.dumps(j.get('requirements',[]),ensure_ascii=False);row['sources']=json.dumps(j['urls'])
  if mode(s)=='Discovery':row['fit']=None;row['fit_reason']=None
  row['recruiting_summary']=('来源冲突，待核对' if status(j)[1]=='Conflict' else ('官网公开招聘' if status(j)==('open','Verified') else ('已关闭' if status(j)[0]=='closed' else '待核实，尚未确认在招')))
  row['candidate_status']='暂不推荐' if not eligible(s,j) else ('可优先考虑' if status(j)==('open','Verified') else '可考虑，招聘状态待核实')
  row['unknown_items']='、'.join(dict.fromkeys(r['key'] for r in j.get('requirements',[]) if r.get('assessment','Unknown')=='Unknown')) or '暂无已识别的待确认项'
  rows.append(row)
 return {'All Jobs':rows,'Recommended':[r for j in recommend(s) for r in rows if r['id']==j['id']], 'Application Tracker':[{k:r[k] for k in ('id','url','company','title','application_status')} for r in rows], 'Verification Log':s['log']}

def extract(src,dest):
 p=Path(src)
 if p.suffix.lower()=='.pdf':
  from pypdf import PdfReader
  text='\n'.join(page.extract_text() or '' for page in PdfReader(p).pages)
 elif p.suffix.lower()=='.docx':
  from docx import Document
  d=Document(p);text='\n'.join([x.text for x in d.paragraphs]+[' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows])
 elif p.suffix.lower() in {'.txt','.md'}:text=p.read_text()
 else:raise ValueError('Supported resume types: txt md pdf docx')
 if not text.strip():raise ValueError('No extracted text: OCR or user review needed')
 Path(dest).write_text(text);return {'characters':len(text),'semantic_profile':'requires Agent review'}

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--state',default='jobmatch-private/state.json');subs=ap.add_subparsers(dest='command',required=True)
 for c in ('init','sources','clarify','intake','coverage','search-plan'):subs.add_parser(c)
 p=subs.add_parser('search-policy');p.add_argument('input')
 p=subs.add_parser('search-log');p.add_argument('record')
 p=subs.add_parser('ingest');p.add_argument('batch')
 p=subs.add_parser('recommend');p.add_argument('--sort',choices=['relevance','match','latest'],default='relevance');p.add_argument('--limit',type=int,default=10);p.add_argument('--include-unverified',action='store_true',help='Explicit exploratory candidate view, not confirmed-open recommendations')
 p=subs.add_parser('status');p.add_argument('id');p.add_argument('value',choices=sorted(APPLICATIONS))
 p=subs.add_parser('source');p.add_argument('id');p.add_argument('url');p.add_argument('--kind',choices=['official','third_party'],required=True);p.add_argument('--enabled',choices=['true','false'],required=True)
 p=subs.add_parser('question');p.add_argument('key');p.add_argument('--result',choices=['asked','skipped','answered'],required=True)
 p=subs.add_parser('export-data');p.add_argument('output')
 p=subs.add_parser('import-tracker');p.add_argument('input')
 p=subs.add_parser('resolve');p.add_argument('id');p.add_argument('--choice',choices=['local','incoming'],required=True)
 p=subs.add_parser('extract');p.add_argument('input');p.add_argument('output')
 a=ap.parse_args();path=Path(a.state);path.parent.mkdir(parents=True,exist_ok=True);lock=Path(str(path)+'.lock')
 mut=a.command in {'init','ingest','status','source','question','import-tracker','resolve','search-log','search-policy'}
 fd=None
 try:
  if mut:fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
  s=load(path);out={}
  if a.command in {'recommend','clarify','export-data'} and not intake(s)['ready']:
   raise ValueError('必要信息未齐全：'+', '.join(intake(s)['missing_required'])+'；先补齐档案后再搜索、推荐或正式导出。')
  if a.command=='intake':out=intake(s)
  elif a.command=='coverage':out=coverage(s)
  elif a.command=='search-plan':out=search_plan(s)
  elif a.command=='search-policy':
   policy=json.loads(Path(a.input).read_text())
   if not isinstance(policy.get('priority_groups'),list) or not all(isinstance(x,str) and x.strip() for x in policy['priority_groups']) or not isinstance(policy.get('allow_expansion'),bool):raise ValueError('Invalid search policy')
   s['search_policy']=policy;out=search_plan(s)
  elif a.command=='search-log':
   record=json.loads(Path(a.record).read_text())
   if not all(record.get(k) for k in ('source_id','kind','query','checked_at','outcome','notes')):raise ValueError('Incomplete search log')
   if record['kind'] not in {'official','third_party'} or record['outcome'] not in {'searched','blocked','unavailable'}:raise ValueError('Invalid search log')
   s.setdefault('search_log',[]).append(record);out=record
  elif a.command=='ingest':out=ingest(s,json.loads(Path(a.batch).read_text()))
  elif a.command=='sources':out={**DEFAULTS,**s['source_overrides']}
  elif a.command=='source':urlkey(a.url);s['source_overrides'][a.id]={'url':a.url,'kind':a.kind,'enabled':a.enabled=='true'};out=s['source_overrides']
  elif a.command=='recommend':out=recommend(s,a.sort,a.limit,a.include_unverified)
  elif a.command=='clarify':out=clarify(s)
  elif a.command=='question':s['questions'][a.key]={'result':a.result,'at':now()}
  elif a.command=='status':update_status(s,a.id,a.value)
  elif a.command=='extract':out=extract(a.input,a.output)
  elif a.command=='export-data':Path(a.output).write_text(json.dumps(export_data(s),ensure_ascii=False,indent=2));out={'saved':a.output}
  elif a.command=='import-tracker':
   if Path(a.input).suffix=='.xlsx':
    from openpyxl import load_workbook
    wb=load_workbook(a.input,read_only=True,data_only=True);rs=list(wb['Application Tracker'].values);rows=[dict(zip(rs[0],r)) for r in rs[1:] if any(r)]
   else:rows=json.loads(Path(a.input).read_text())
   out=import_tracker(s,rows)
  elif a.command=='resolve':
   c=next(c for c in s['conflicts'] if c['id']==a.id and not c['resolved']);j=next(j for j in s['jobs'] if j['id']==c['job_id'])
   if j['application_status']!=c['local']:raise ValueError('Local status changed; reimport tracker')
   if a.choice=='incoming':update_status(s,c['job_id'],c['incoming'])
   c['resolved']=True
  if mut:
   with tempfile.NamedTemporaryFile(mode='w',dir=path.parent,delete=False) as f:json.dump(s,f,ensure_ascii=False,indent=2);temp=f.name
   os.chmod(temp,0o600);os.replace(temp,path)
  print(json.dumps(out,ensure_ascii=False,indent=2))
 finally:
  if fd is not None:os.close(fd);lock.unlink(missing_ok=True)
if __name__=='__main__':main()
